"""The model as JSON, for the UI and other tools (PRD #10, docs/DESIGN.md "JSON export"), and the progress table as
CSV or Excel (docs/features/export-csv-xlsx.md)."""
import csv
import io
import json
from dataclasses import asdict
from datetime import datetime, timedelta, timezone

from .model import View
from .rules import ALL

SCHEMA = 2  # bump on breaking changes, so readers can tell what they got. 2: one catalog tree, progress per slot


def to_dict(lib, mods=None):
    return {
        "schema": SCHEMA,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "mods_folder": str(mods.source) if mods and mods.source else None,
        **asdict(lib),
    }


def to_json(lib, mods=None):
    return json.dumps(to_dict(lib, mods), indent=2, ensure_ascii=False, default=str)


# ------------------------------------------------------------------ the table, one row per mod

TICKS = 10_000_000
# The page's three statuses (app.js PAGE_STATUS, docs/UI.md "Statuses"): a file is read by people, like the page.
PAGE_STATUS = {"started": "in progress", "complete": "completed", "all opened done": "completed",
               "not opened": "not started"}


def _duration(ticks):
    return timedelta(seconds=ticks // TICKS) if ticks else None


def _view(mod, key):
    """The mod's view in this slot, or, for a slot it wasn't played in, not started with the catalog's totals (as
    the page's viewOf)."""
    v = mod.progress.get(key)
    return v if v is not None else View(status="not started", sides_total=mod.progress[ALL].sides_total)


def _counted_sides(ch, key):
    """The sides the page counts in this slot (app.js countedSides): the ones that exist, or were opened there."""
    return [s for s in ch.sides.values() if s.exists or key in s.progress]


def _best(mod, key):
    """The best time of the mod's only side, as the page's Best column shows it on a mod row; None otherwise."""
    chapters = mod.chapters()
    if len(chapters) != 1:
        return None
    sides = _counted_sides(chapters[0], key)
    if len(sides) != 1 or key not in sides[0].progress:
        return None
    return _duration(sides[0].progress[key].best_ticks)


def _checkpoint(mod, v, key):
    """The checkpoint's name and room; on a mod with several chapters, which chapter and side ("Ridge, B side:
    Start (b-00)"), or with several sides, which side, as the page shows under it."""
    cp = v.latest_checkpoint
    if not cp:
        return ""
    text = f"{cp['title']} ({cp['room']})" if cp["title"] else cp["room"]
    chapters = mod.chapters()
    if len(chapters) > 1:
        ch = mod.find_chapter(cp["sid"])
        return f"{(ch.title if ch else '') or cp['sid'].rpartition('/')[2]}, {cp['side']} side: {text}"
    if chapters and len(_counted_sides(chapters[0], key)) > 1:
        return f"{cp['side']} side: {text}"
    return text


def _played(v):
    return v.status not in ("not started", "not opened")


# id -> (header, value(mod, view, slot key)). The page's Export panel lists the same IDs (app.js EXPORT_COLUMNS);
# a file's columns come in this order, whatever order they were asked in.
COLUMNS = {
    "mod": ("Mod", lambda m, v, k: m.name),
    "id": ("Mod ID", lambda m, v, k: m.id),
    "author": ("Author", lambda m, v, k: m.author),
    "status": ("Status", lambda m, v, k: PAGE_STATUS.get(v.status, v.status)),
    "sides_done": ("Sides cleared", lambda m, v, k: v.sides_done),
    "sides_total": ("Sides total", lambda m, v, k: v.sides_total),
    "deaths": ("Deaths", lambda m, v, k: v.deaths if _played(v) else None),
    "time": ("Time", lambda m, v, k: _duration(v.ticks) if _played(v) else None),
    "best": ("Best", lambda m, v, k: _best(m, k)),
    "berries": ("Berries", lambda m, v, k: v.berries if _played(v) else None),
    "checkpoint": ("Checkpoint", lambda m, v, k: _checkpoint(m, v, k)),
    "rating": ("Rating", lambda m, v, k: m.user.get("rating") or None),
    "difficulty": ("Difficulty", lambda m, v, k: m.user.get("difficulty", "")),
    "tags": ("Tags", lambda m, v, k: "; ".join(m.user.get("tags") or [])),
    "note": ("Note", lambda m, v, k: m.user.get("note", "")),
    "slot": ("Slot", lambda m, v, k: int(v.slot) if v.slot and v.slot.isdigit() else v.slot),  # All slots only
}


def to_rows(lib, key, cols, ids):
    """The header, then one row per mod ID, in the order given. key: a slot key or "all". Raises ValueError on an
    unknown slot, column or mod ID."""
    if key != ALL and not any(s.key == key for s in lib.slots):
        raise ValueError(f"no slot {key!r}")
    bad = [c for c in cols if c not in COLUMNS or (c == "slot" and key != ALL)]
    if bad or not cols:
        raise ValueError(f"bad columns: {bad or 'none'}")
    by_id = {m.id: m for m in lib.mods}
    missing = [i for i in ids if i not in by_id]
    if missing:
        raise ValueError(f"no mod {missing[0]!r}")
    cols = [c for c in COLUMNS if c in cols]
    rows = [[COLUMNS[c][0] for c in cols]]
    for i in dict.fromkeys(ids):  # once each
        m = by_id[i]
        v = _view(m, key)
        rows.append([COLUMNS[c][1](m, v, key) for c in cols])
    return rows


def _text(value):
    if isinstance(value, timedelta):
        s = int(value.total_seconds())
        return f"{s // 3600}:{s // 60 % 60:02}:{s % 60:02}"
    return "" if value is None else value


def _safe(value):
    """Text a spreadsheet would read as a formula (a note starting with "="), kept as text."""
    return "'" + value if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def to_csv(rows):
    """UTF-8 with a BOM, so Excel shows accented names right; Excel's own \\r\\n line ends."""
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    for row in rows:
        w.writerow([_safe(_text(x)) for x in row])
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(rows):
    """One sheet: a bold, frozen header with a filter; numbers as numbers, times as durations."""
    from openpyxl import Workbook  # only when asked: it takes ~0.2 s to import
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Progress"
    widths = [0] * len(rows[0])
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell = ws.cell(r, c, value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula, even when a note starts with "="
            elif isinstance(value, timedelta):
                cell.number_format = "[h]:mm:ss"
            if r == 1:
                cell.font = Font(bold=True)
            longest = max((len(x) for x in str(_text(value)).splitlines()), default=0)
            widths[c - 1] = max(widths[c - 1], longest)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
