"""The table as CSV or Excel (docs/features/export-csv-xlsx.md), on the made-up saves."""
import csv
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from celeste_tracker.core import load_library
from celeste_tracker.export import COLUMNS, to_csv, to_rows, to_xlsx

from conftest import SLOTS


@pytest.fixture
def lib(mods):
    lib = load_library([(1, SLOTS / "1.celeste"), (2, SLOTS / "2.celeste")], mods)
    sides = next(m for m in lib.mods if m.id == "Sides Mod")
    sides.user = {"rating": 4, "tags": ["golden", "fun"], "note": "=1+1, \"quoted\"\nsecond line"}
    return lib


def ids(lib):
    return [m.id for m in lib.mods]


def test_header_then_one_row_per_mod_in_the_order_given(lib):
    order = list(reversed(ids(lib)))
    rows = to_rows(lib, "all", ["mod", "id"], order + order[:1])  # a repeated ID comes once
    assert rows[0] == ["Mod", "Mod ID"]
    assert [r[1] for r in rows[1:]] == order


def test_columns_come_in_a_fixed_order(lib):
    assert to_rows(lib, "all", ["note", "mod", "time"], [])[0] == ["Mod", "Time", "Note"]


def test_values(lib):
    rows = to_rows(lib, "all", list(COLUMNS), ["Sides Mod", "Celeste"])
    head, sides, vanilla = rows
    row = dict(zip(head, sides))
    assert row["Mod"] == "The Forest" and row["Status"] == "in progress"
    assert (row["Sides cleared"], row["Sides total"]) == (2, 3)
    assert isinstance(row["Time"], timedelta) and row["Deaths"] > 0
    assert row["Rating"] == 4 and row["Tags"] == "golden; fun" and row["Difficulty"] == ""
    assert row["Checkpoint"].startswith("C side: ") and "c-01" in row["Checkpoint"]  # one chapter, three sides
    assert isinstance(row["Slot"], int)
    assert dict(zip(head, vanilla))["Rating"] is None


def test_not_played_in_a_slot(lib):
    # Blizzard isn't in slot 2: not started, with the catalog's totals, no numbers
    played = {m.id for m in lib.mods if "2" in m.progress}
    other = next(i for i in ids(lib) if i not in played)
    row = to_rows(lib, "2", ["status", "sides_total", "deaths", "time"], [other])[1]
    assert row[0] == "not started" and row[2] is None and row[3] is None


def test_bad_input(lib):
    with pytest.raises(ValueError):
        to_rows(lib, "9", ["mod"], [])
    with pytest.raises(ValueError):
        to_rows(lib, "all", ["nope"], [])
    with pytest.raises(ValueError):
        to_rows(lib, "1", ["slot"], [])  # Slot is only for All slots
    with pytest.raises(ValueError):
        to_rows(lib, "all", [], [])
    with pytest.raises(ValueError):
        to_rows(lib, "all", ["mod"], ["Not A Mod"])


def test_csv(lib):
    data = to_csv(to_rows(lib, "all", ["mod", "time", "note"], ["Sides Mod"]))
    assert data.startswith(b"\xef\xbb\xbf") and b"\r\n" in data
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))
    assert rows[0] == ["Mod", "Time", "Note"]
    name, time, note = rows[1]
    assert name == "The Forest" and time.count(":") == 2
    assert note == "'=1+1, \"quoted\"\nsecond line"  # kept whole, and never a formula


def test_xlsx(lib):
    data = to_xlsx(to_rows(lib, "all", ["mod", "sides_done", "time", "note"], ["Sides Mod"]))
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["Mod", "Sides cleared", "Time", "Note"]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D2"
    assert ws["B2"].value == 2
    assert ws["C2"].number_format == "[h]:mm:ss" and isinstance(ws["C2"].value, timedelta)
    assert ws["D2"].data_type == "s" and ws["D2"].value.startswith("=1+1")  # text, not a formula
